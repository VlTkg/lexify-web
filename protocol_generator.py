import io
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex: str):
    """Устанавливает цвет фона ячейки таблицы."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Настраивает внутренние отступы (padding) ячейки."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_disagreements_protocol(analysis_data: dict, contract_title: str = "Договору оказания услуг") -> io.BytesIO:
    """
    Принимает JSON-структуру анализа договора и генерирует .docx файл
    Протокола Разногласий. Возвращает io.BytesIO поток.
    """
    doc = Document()

    # Поля страницы (2 см со всех сторон)
    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    # Настройка базового стиля текста
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # --- ЗАГОЛОВОК И ШАПКА ---
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("ПРОТОКОЛ РАЗНОГЛАСИЙ")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"к {contract_title}\n")
    sub_run.font.size = Pt(12)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    # --- ТАБЛИЦА РАЗНОГЛАСИЙ ---
    # Фильтруем только критические (red) и предупреждающие (yellow) риски
    risks = [r for r in analysis_data.get("risks", []) if r.get("level") in ["red", "yellow"]]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    col_widths = [Inches(1.0), Inches(2.2), Inches(2.3), Inches(1.8)]

    # Заполняем шапку таблицы
    hdr_cells = table.rows[0].cells
    headers = ["Пункт", "Редакция Заказчика", "Предлагаемая редакция Исполнителя", "Обоснование"]

    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_background(cell, "1F2937") # Тёмный современный цвет шапки
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header_text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Заполнение содержимым
    if not risks:
        row_cells = table.add_row().cells
        for i in range(4):
            row_cells[i].width = col_widths[i]
        row_cells[0].merge(row_cells[3])
        row_cells[0].text = "Существенных рисков и замечаний к редакции договора не выявлено."
    else:
        for risk in risks:
            row_cells = table.add_row().cells

            for i in range(4):
                row_cells[i].width = col_widths[i]
                set_cell_margins(row_cells[i], top=100, bottom=100, left=100, right=100)

            # 1. Пункт
            row_cells[0].text = risk.get("clause_num", "—")
            row_cells[0].paragraphs[0].runs[0].bold = True

            # 2. Описание риска / Редакция заказчика
            row_cells[1].text = risk.get("description", "")

            # 3. Предложенная правка
            row_cells[2].text = risk.get("suggested_fix", "")

            # 4. Обоснование
            justification = risk.get("reasoning") or "Приведение условий договора к балансу интересов сторон и стандартной рыночной практике."
            row_cells[3].text = justification

    doc.add_paragraph("\n")

    # --- БЛОК ПОДПИСЕЙ ---
    doc.add_paragraph("Настоящий Протокол разногласий является неотъемлемой частью Договора.\n")

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell_left = sig_table.rows[0].cells[0]
    cell_right = sig_table.rows[0].cells[1]

    cell_left.width = Inches(3.6)
    cell_right.width = Inches(3.6)

    cell_left.text = "ПОДРЯДЧИК / ИСПОЛНИТЕЛЬ:\n\n\n_______________________ / ____________ /\nМ.П."
    cell_right.text = "ЗАКАЗЧИК:\n\n\n_______________________ / ____________ /\nМ.П."

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_stream.seek(0)
    return file_stream
