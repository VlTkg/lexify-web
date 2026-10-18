import os
import io
import httpx
from fastapi import FastAPI, HTTPException, UploadFile, File, Form, Request
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel
from typing import Optional, List
from google import genai
from pypdf import PdfReader
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Импорты для работы с SQLAlchemy и PostgreSQL
from sqlalchemy import create_engine, Column, BigInteger, Integer, String, Text, DateTime, Boolean
from sqlalchemy.ext.declarative import declarative_base
from sqlalchemy.orm import sessionmaker
from datetime import datetime

app = FastAPI()

# Монтируем статические файлы
app.mount("/static", StaticFiles(directory="static"), name="static")

# Переменные окружения
FEEDBACK_WEBHOOK_URL = os.environ.get("GOOGLE_SHEET_WEBHOOK_URL", "")
GEMINI_API_KEY = os.environ.get("GEMINI_API_KEY", "")
DATABASE_URL = os.environ.get("DATABASE_URL", "")
BOT_TOKEN = os.environ.get("BOT_TOKEN", "") # Токен вашего Telegram бота от BotFather

# Список Telegram ID администраторов / разработчиков с безлимитными проверками
ADMIN_IDS = [8804083078]

# Настройка подключения к PostgreSQL через SQLAlchemy
if DATABASE_URL and DATABASE_URL.startswith("postgres://"):
    DATABASE_URL = DATABASE_URL.replace("postgres://", "postgresql://", 1)

engine = None
SessionLocal = None
Base = declarative_base()

# Модель пользователя в базе данных
class UserDB(Base):
    __tablename__ = "users"

    telegram_id = Column(BigInteger, primary_key=True, index=True)
    first_name = Column(String, default="")
    username = Column(String, default="")
    free_checks_left = Column(Integer, default=3)
    is_pro = Column(Boolean, default=False)
    created_at = Column(DateTime, default=datetime.utcnow)

# Модель логов анализа
class AnalysisLogDB(Base):
    __tablename__ = "analysis_logs"

    id = Column(Integer, primary_key=True, index=True)
    telegram_id = Column(BigInteger, nullable=True)
    doc_type = Column(String, default="")
    jurisdiction = Column(String, default="")
    risk_level = Column(String, default="")
    created_at = Column(DateTime, default=datetime.utcnow)

# Модель для логирования успешных платежей (защита от повторных начислений)
class PaymentLogDB(Base):
    __tablename__ = "payment_logs"

    id = Column(Integer, primary_key=True, index=True)
    telegram_id = Column(BigInteger, nullable=False)
    charge_id = Column(String, unique=True, index=True, nullable=False)
    payload = Column(String, default="")
    created_at = Column(DateTime, default=datetime.utcnow)

if DATABASE_URL:
    try:
        engine = create_engine(DATABASE_URL)
        SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
        Base.metadata.create_all(bind=engine)
        print("База данных подключена, таблицы успешно инициализированы.")
    except Exception as e:
        print(f"Ошибка подключения к базе данных: {e}")

# Инициализация клиента Gemini
ai_client = None
if GEMINI_API_KEY:
    ai_client = genai.Client(api_key=GEMINI_API_KEY)

class FeedbackSchema(BaseModel):
    type: str = ""
    rating: int = 0
    sentiment: str = ""
    feedback: str = ""
    doc_type: str = ""
    jurisdiction: str = ""

class AnalyzeSchema(BaseModel):
    text: str
    doc_type: str = "Общий"
    jurisdiction: str = "Россия (РФ)"
    telegram_id: Optional[int] = None
    first_name: Optional[str] = ""
    username: Optional[str] = ""

class InvoiceRequest(BaseModel):
    telegram_id: int
    package_type: str = "5_checks" # Варианты: '5_checks' или '15_checks'

# Модели для структурированного ответа анализа рисков
class RiskItem(BaseModel):
    category: str  # 'red', 'yellow', 'green'
    title: str     # Краткий заголовок риска
    quote: str     # Цитата или упоминание пункта договора
    explanation: str  # В чем заключается риск для фрилансера/исполнителя
    recommendation: str  # Готовый текст правки или рекомендация для общения с заказчиком

class AnalysisResponse(BaseModel):
    overall_risk_level: str  # 'Низкий', 'Средний', 'Высокий'
    summary: str  # Краткое общее резюме (1-2 предложения)
    risks: List[RiskItem]

# Модели для генерации Протокола разногласий в docx
class ProtocolItem(BaseModel):
    title: str
    quote: str
    explanation: str
    recommendation: str

class ProtocolRequest(BaseModel):
    doc_type: str = "Общий"
    jurisdiction: str = "Россия (РФ)"
    telegram_id: Optional[int] = None
    risks: List[ProtocolItem]

@app.get("/")
async def serve_index():
    return FileResponse("static/index.html")

# Эндпоинт для получения информации о пользователе и остатке лимита
@app.get("/user-info/{telegram_id}")
async def get_user_info(telegram_id: int):
    if telegram_id in ADMIN_IDS:
        return {"telegram_id": telegram_id, "free_checks_left": 999, "is_pro": True}

    if not SessionLocal:
        return {"telegram_id": telegram_id, "free_checks_left": 3, "is_pro": False}
    
    db = SessionLocal()
    try:
        user = db.query(UserDB).filter(UserDB.telegram_id == telegram_id).first()
        if not user:
            user = UserDB(telegram_id=telegram_id, free_checks_left=3, is_pro=False)
            db.add(user)
            db.commit()
            db.refresh(user)
        return {
            "telegram_id": user.telegram_id,
            "free_checks_left": user.free_checks_left,
            "is_pro": user.is_pro
        }
    finally:
        db.close()

# Эндпоинт создания инвойса оплаты через Telegram Stars
@app.post("/create-invoice")
async def create_invoice(data: InvoiceRequest):
    if not BOT_TOKEN:
        raise HTTPException(status_code=500, detail="BOT_TOKEN не настроен на сервере.")
    
    packages = {
        "5_checks": {"title": "Пакет 5 проверок", "description": "5 дополнительных экспресс-анализов договоров", "amount": 50},
        "15_checks": {"title": "Пакет 15 проверок", "description": "15 дополнительных экспресс-анализов договоров", "amount": 120}
    }
    
    pkg = packages.get(data.package_type, packages["5_checks"])
    
    url = f"https://api.telegram.org/bot{BOT_TOKEN}/createInvoiceLink"
    payload = {
        "title": pkg["title"],
        "description": pkg["description"],
        "payload": f"pack_{data.package_type}_{data.telegram_id}",
        "currency": "XTR",
        "prices": [{"label": pkg["title"], "amount": pkg["amount"]}]
    }
    
    async with httpx.AsyncClient() as client:
        response = await client.post(url, json=payload)
        res_data = response.json()
        if res_data.get("ok"):
            return {"invoice_link": res_data["result"]}
        else:
            raise HTTPException(status_code=500, detail=f"Ошибка создания счета Telegram: {res_data.get('description')}")

# Webhook для получения обновлений от Telegram (обработка успешных платежей с защитой от дубликатов)
@app.post("/telegram-webhook")
async def telegram_webhook(request: Request):
    body = await request.json()
    
    if "message" in body and "successful_payment" in body["message"]:
        payment = body["message"]["successful_payment"]
        payload = payment.get("invoice_payload", "")
        user_id = body["message"]["from"]["id"]
        charge_id = payment.get("telegram_payment_charge_id", "")
        
        if payload.startswith("pack_"):
            parts = payload.split("_")
            if len(parts) >= 3:
                checks_to_add = int(parts[1]) if parts[1].isdigit() else 5
                
                if SessionLocal:
                    db = SessionLocal()
                    try:
                        existing_payment = db.query(PaymentLogDB).filter(PaymentLogDB.charge_id == charge_id).first()
                        if existing_payment:
                            print(f"Платеж {charge_id} уже был обработан ранее.")
                            return {"status": "ok"}
                        
                        user = db.query(UserDB).filter(UserDB.telegram_id == user_id).first()
                        if user:
                            user.free_checks_left += checks_to_add
                            
                            new_payment_log = PaymentLogDB(
                                telegram_id=user_id,
                                charge_id=charge_id,
                                payload=payload
                            )
                            db.add(new_payment_log)
                            db.commit()
                            print(f"Пользователю {user_id} успешно начислено {checks_to_add} проверок за Telegram Stars (Charge ID: {charge_id})")
                    except Exception as e:
                        db.rollback()
                        print(f"Ошибка начисления проверок по платежу: {e}")
                    finally:
                        db.close()
                        
    return {"status": "ok"}

# Вспомогательные функции форматирования для красивого генератора .docx
def set_cell_background(cell, fill_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

# Эндпоинт генерации Протокола разногласий в профессиональном стиле .docx
@app.post("/generate-protocol")
async def generate_protocol(data: ProtocolRequest):
    doc = docx.Document()

    # Поля страницы
    for section in doc.sections:
        section.top_margin = Inches(0.78)
        section.bottom_margin = Inches(0.78)
        section.left_margin = Inches(0.78)
        section.right_margin = Inches(0.78)

    # Базовый стиль текста
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # Заголовок
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("ПРОТОКОЛ РАЗНОГЛАСИЙ")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor(0x11, 0x18, 0x27)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run(f"к проекту договора ({data.doc_type}), юрисдикция: {data.jurisdiction}\n")
    sub_run.font.size = Pt(12)
    sub_run.italic = True
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    date_p = doc.add_paragraph()
    date_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_run = date_p.add_run(f"Дата формирования: {datetime.now().strftime('%d.%m.%Y')}\n")
    date_run.font.size = Pt(10)

    # Таблица разногласий
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'

    col_widths = [Inches(1.0), Inches(2.2), Inches(2.3), Inches(1.8)]

    hdr_cells = table.rows[0].cells
    headers = ["Пункт", "Редакция Заказчика", "Предлагаемая редакция Исполнителя", "Обоснование"]

    for i, header_text in enumerate(headers):
        cell = hdr_cells[i]
        cell.width = col_widths[i]
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_background(cell, "1F2937")
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = p.add_run(header_text)
        run.bold = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    if not data.risks:
        row_cells = table.add_row().cells
        for i in range(4):
            row_cells[i].width = col_widths[i]
        row_cells[0].merge(row_cells[3])
        row_cells[0].text = "Существенных рисков и замечаний к редакции договора не выявлено."
    else:
        for risk in data.risks:
            row_cells = table.add_row().cells

            for i in range(4):
                row_cells[i].width = col_widths[i]
                set_cell_margins(row_cells[i], top=100, bottom=100, left=100, right=100)

            row_cells[0].text = risk.title or "—"
            row_cells[0].paragraphs[0].runs[0].bold = True
            row_cells[1].text = risk.quote or ""
            row_cells[2].text = risk.recommendation or ""
            row_cells[3].text = risk.explanation or "Защита интересов исполнителя"

    doc.add_paragraph("\n")
    doc.add_paragraph("Настоящий Протокол разногласий является неотъемлемой частью Договора.\n")

    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER

    cell_left = sig_table.rows[0].cells[0]
    cell_right = sig_table.rows[0].cells[1]

    cell_left.width = Inches(3.6)
    cell_right.width = Inches(3.6)

    cell_left.text = "ПОДРЯДЧИК / ИСПОЛНИТЕЛЬ:\n\n\n_______________________ / ____________ /\nМ.П."
    cell_right.text = "ЗАКАЗЧИК:\n\n\n_______________________ / ____________ /\nМ.П."

    file_stream = io.BytesIO()
    doc.save(file_stream)
    file_bytes = file_stream.getvalue()
    file_stream.seek(0)

    # Если передан telegram_id и настроен BOT_TOKEN — отправляем файл прямо в чат Telegram
    if data.telegram_id and BOT_TOKEN:
        try:
            async with httpx.AsyncClient() as client:
                files = {
                    'document': ('Protocol_Raznogasiy.docx', file_bytes, 'application/vnd.openxmlformats-officedocument.wordprocessingml.document')
                }
                payload = {
                    'chat_id': data.telegram_id,
                    'caption': '📄 Ваш Протокол разногласий готов!'
                }
                tg_resp = await client.post(
                    f"https://api.telegram.org/bot{BOT_TOKEN}/sendDocument",
                    data=payload,
                    files=files,
                    timeout=15.0
                )
                if tg_resp.status_code == 200:
                    return {"status": "ok", "message": "Файл успешно отправлен в чат с ботом!"}
        except Exception as e:
            print(f"Ошибка отправки файла в Telegram: {e}")

    # Фолбэк для браузера (если вызов вне Telegram)
    return StreamingResponse(
        io.BytesIO(file_bytes),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": "attachment; filename=Protocol_Raznogasiy.docx"}
    )

@app.post("/feedback")
async def receive_feedback(data: FeedbackSchema):
    if not FEEDBACK_WEBHOOK_URL:
        print(f"Feedback received (no webhook configured): {data}")
        return {"status": "ok", "message": "Logged locally"}

    try:
        async with httpx.AsyncClient() as client:
            response = await client.post(
                FEEDBACK_WEBHOOK_URL,
                json=data.dict(),
                timeout=10.0
            )
            if response.status_code == 200:
                return {"status": "ok", "message": "Feedback saved to Google Sheet"}
            else:
                raise HTTPException(status_code=500, detail="Failed to post to Google Sheet")
    except Exception as e:
        print(f"Error sending feedback: {e}")
        return {"status": "error", "detail": str(e)}

@app.post("/analyze")
async def analyze_contract(data: AnalyzeSchema):
    if not ai_client:
        raise HTTPException(status_code=500, detail="GEMINI_API_KEY не настроен на сервере.")

    contract_content = data.text.strip()
    if not contract_content:
        raise HTTPException(status_code=400, detail="Текст договора пуст.")

    checks_left = 3
    if data.telegram_id and SessionLocal:
        db = SessionLocal()
        try:
            user = db.query(UserDB).filter(UserDB.telegram_id == int(data.telegram_id)).first()
            if not user:
                user = UserDB(
                    telegram_id=int(data.telegram_id),
                    first_name=data.first_name or "",
                    username=data.username or "",
                    free_checks_left=3
                )
                db.add(user)
                db.commit()
                db.refresh(user)
            
            # Проверка администратора или PRO-статуса
            if int(data.telegram_id) in ADMIN_IDS or user.is_pro:
                checks_left = 999  # Безлимитные проверки для администратора
            else:
                if user.free_checks_left <= 0:
                    raise HTTPException(
                        status_code=403, 
                        detail="У вас закончились бесплатные проверки. Пополните баланс звездами Telegram!"
                    )
                user.free_checks_left -= 1
                db.commit()
                db.refresh(user)
                checks_left = user.free_checks_left

        except HTTPException as he:
            raise he
        except Exception as db_err:
            print(f"Ошибка работы с БД пользователей: {db_err}")
        finally:
            db.close()

    prompt = f"""
Ты — профессиональный юрист-аналитик, специализирующийся на защите интересов фрилансеров, подрядчиков и специалистов.
Проведи экспресс-анализ представленного договора и выяви потенциальные риски для Исполнителя.
Пиши современным, понятным и деловым языком. Избегай архаичных юридических канцеляризмов (не используй слова вроде «кабальный», заменяй их на «неприемлемые условия», «критические риски» или «односторонние невыгодные ограничения»).

Контекст анализа:
- Тип договора / Занятость: {data.doc_type}
- Юрисдикция: {data.jurisdiction}

Текст договора:
---
{contract_content[:15000]}
---

Сформируй разбор строго в формате JSON:
1. `overall_risk_level`: Оцени общий уровень риска (строго одно из значений: 'Низкий', 'Средний', 'Высокий').
2. `summary`: Краткий вывод по договору в 1-2 предложениях (без канцелярита).
3. `risks`: Список выявленных пунктов и рисков (3-6 пунктов). Для каждого укажи:
   - `category`: Цветовой маркер риска ('red' — критический, 'yellow' — средний/внимание, 'green' — безопасное/выгодное условие).
   - `title`: Четкое название риска (например, "Неограниченный штраф за задержку", "Передача прав до оплаты").
   - `quote`: Цитата или фрагмент пункта договора (если применимо).
   - `explanation`: Чем именно это условие грозит Исполнителю.
   - `recommendation`: Готовая формулировка правки, которую фрилансер может скопировать и выслать заказчику для защиты своих интересов.
"""

    try:
        response = ai_client.models.generate_content(
            model="gemini-3.8-flash",
            contents=prompt,
            config={
                "response_mime_type": "application/json",
                "response_schema": AnalysisResponse,
            }
        )
        
        if SessionLocal:
            try:
                db = SessionLocal()
                db_log = AnalysisLogDB(
                    telegram_id=int(data.telegram_id) if data.telegram_id else None,
                    doc_type=data.doc_type,
                    jurisdiction=data.jurisdiction,
                    risk_level="Анализ выполнен"
                )
                db.add(db_log)
                db.commit()
                db.close()
            except Exception as db_err:
                print(f"Ошибка сохранения лога в БД: {db_err}")

        return {
            "status": "ok", 
            "result": response.text,
            "checks_left": checks_left
        }
    except HTTPException as he:
        raise he
    except Exception as e:
        print(f"Gemini API error: {e}")
        raise HTTPException(status_code=500, detail=f"Ошибка анализа через Gemini API: {str(e)}")
